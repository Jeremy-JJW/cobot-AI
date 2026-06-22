"""將使用說明書 markdown 轉換為格式規整的 DOCX。"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def _set_cell_shading(cell, color: str):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def _add_table(doc, rows):
    """Add a formatted table from parsed markdown rows."""
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text.strip()
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = 'Microsoft JhengHei'
                    if i == 0:  # Header row
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            if i == 0:
                _set_cell_shading(cell, "2F5496")
            elif i % 2 == 0:
                _set_cell_shading(cell, "D6E4F0")

    doc.add_paragraph()  # spacing after table


def md_to_docx(md_path: str, docx_path: str):
    md = Path(md_path).read_text(encoding="utf-8")
    doc = Document()

    # -- style defaults --
    style = doc.styles['Normal']
    style.font.name = 'Microsoft JhengHei'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

    for level in range(1, 4):
        hs = doc.styles[f'Heading {level}']
        hs.font.name = 'Microsoft JhengHei'
        hs.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        if level == 1:
            hs.font.size = Pt(22)
        elif level == 2:
            hs.font.size = Pt(16)
        else:
            hs.font.size = Pt(13)

    lines = md.split('\n')
    in_table = False
    table_rows: list[list[str]] = []
    in_code_block = False
    code_lines: list[str] = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run('\n'.join(code_lines))
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
        # grey background via shading
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="#F2F2F2"/>')
        p._element.get_or_add_pPr().append(shading)
        code_lines = []

    for raw_line in lines:
        stripped = raw_line.strip()

        # -- code block toggle --
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
                in_code_block = False
            else:
                flush_code()
                in_code_block = True
            continue
        if in_code_block:
            code_lines.append(raw_line)
            continue

        flush_code()

        # -- skip standalone separators (---) --
        if stripped == '---':
            continue

        # -- headings --
        hm = re.match(r'^(#{1,3})\s+(.+)$', stripped)
        if hm:
            level = len(hm.group(1))
            doc.add_heading(hm.group(2), level=level)
            continue

        # -- table rows --
        if '|' in stripped and stripped.startswith('|'):
            cells = [c.strip() for c in stripped.split('|') if c.strip() != '']
            if re.match(r'^[\s\|:\-]+$', stripped):
                continue  # skip separator row
            in_table = True
            table_rows.append(cells)
            continue
        else:
            if in_table and table_rows:
                _add_table(doc, table_rows)
                table_rows = []
                in_table = False

        # -- bullet list --
        bm = re.match(r'^(\s*)[\-\*]\s+(.+)$', stripped)
        if bm:
            p = doc.add_paragraph(bm.group(2), style='List Bullet')
            continue

        # -- numbered list --
        nm = re.match(r'^(\s*)\d+\.\s+(.+)$', stripped)
        if nm:
            p = doc.add_paragraph(nm.group(2), style='List Number')
            continue

        # -- blockquote --
        if stripped.startswith('>'):
            text = re.sub(r'^>\s*', '', stripped)
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Inches(0.3)
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            continue

        # -- bold + plain paragraph (handle **bold** inline) --
        if stripped:
            p = doc.add_paragraph()
            # Split on **...**
            parts = re.split(r'(\*\*.+?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
            continue

        # -- empty line: spacing --
        doc.add_paragraph()

    # flush remaining
    flush_code()
    if in_table and table_rows:
        _add_table(doc, table_rows)

    # -- cover page title --
    # insert cover at very beginning
    first_para = doc.paragraphs[0]
    cover = first_para.insert_paragraph_before()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(6):
        cover.add_run('\n')
    run = cover.add_run('COBOT AI\n機械臂控制系統\n使用說明書')
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    run.bold = True
    run.font.name = 'Microsoft JhengHei'

    r_sub = cover.add_run('\n\n繁體中文版')
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    r_sub.font.name = 'Microsoft JhengHei'

    # insert page break between cover and first content
    break_p = first_para.insert_paragraph_before()
    break_run = break_p.add_run()
    break_run.add_break(WD_BREAK.PAGE)

    doc.save(docx_path)
    print(f"✓ 已生成: {docx_path}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "使用說明書（台灣版）.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else src.replace(".md", ".docx")
    md_to_docx(src, dst)
